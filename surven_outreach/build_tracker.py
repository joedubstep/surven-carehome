"""Generate the Surven Email Outreach Tracker (xlsx) from the engine DB.

Run:  python -m surven_outreach.build_tracker
Output: data/Surven_Email_Tracker.xlsx  (upload to Drive; refresh via File -> Import -> Replace)
"""
import datetime as dt
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from surven_outreach import db as dbm
from surven_outreach.config import load_config

HEADER_FILL = PatternFill("solid", fgColor="2F3B33")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SECTION = Font(bold=True, size=11, color="2F3B33")
ITALIC = Font(italic=True, color="777777")


def _pct(part, whole):
    return (part / whole) if whole else 0.0


def build(path="data/Surven_Email_Tracker.xlsx"):
    c = load_config()
    con = dbm.connect(c.db_path)

    contacted = con.execute(
        """SELECT business, email, category, city, personalized, first_contact_date,
                  current_step, status, replied_at
           FROM prospects WHERE status NOT IN ('queued','skipped')
           ORDER BY first_contact_date DESC, business"""
    ).fetchall()
    total_sends = con.execute("SELECT COUNT(*) FROM sends").fetchone()[0]
    sends_by_step = {s: 0 for s in range(1, 6)}
    for r in con.execute("SELECT step, COUNT(*) c FROM sends GROUP BY step"):
        sends_by_step[r["step"]] = r["c"]
    bookings = con.execute("SELECT COUNT(*) FROM prospects WHERE status='booked'").fetchone()[0]

    def cnt(status):
        return con.execute("SELECT COUNT(*) FROM prospects WHERE status=?", (status,)).fetchone()[0]

    n_contacted = len(contacted)
    replied, bounced, opted = cnt("replied"), cnt("bounced"), cnt("opted_out")
    active, completed, queued = cnt("active"), cnt("completed"), cnt("queued")

    niche = defaultdict(lambda: [0, 0])     # category -> [contacted, replied]
    pers = {1: [0, 0], 0: [0, 0]}           # personalized -> [contacted, replied]
    for r in contacted:
        is_reply = r["status"] == "replied"
        cat = r["category"] or "(unknown)"
        niche[cat][0] += 1
        niche[cat][1] += int(is_reply)
        pf = 1 if r["personalized"] else 0
        pers[pf][0] += 1
        pers[pf][1] += int(is_reply)

    wb = Workbook()

    # ---------------- Dashboard ----------------
    d = wb.active
    d.title = "Dashboard"
    d["A1"] = "Surven — Email Outreach Tracker"
    d["A1"].font = TITLE_FONT
    d["A2"] = f"updated {dt.date.today().isoformat()} · regenerate from the engine to refresh"
    d["A2"].font = ITALIC

    r = 4

    def section(title):
        nonlocal r
        d.cell(r, 1, title).font = SECTION
        r += 1

    def kv(label, value, is_pct=False):
        nonlocal r
        d.cell(r, 1, label)
        cell = d.cell(r, 2, value)
        if is_pct:
            cell.number_format = "0.0%"
        r += 1

    section("FUNNEL")
    kv("Businesses contacted", n_contacted)
    kv("Emails sent (all touches)", total_sends)
    kv("Replied", replied)
    kv("Appointments booked", bookings)
    kv("Bounced", bounced)
    kv("Opted out", opted)
    kv("In sequence (active)", active)
    kv("Finished, no reply", completed)
    kv("Queue remaining", queued)
    r += 1

    section("RATES")
    kv("Reply rate", _pct(replied, n_contacted), is_pct=True)
    kv("Email -> appointment rate", _pct(bookings, n_contacted), is_pct=True)
    kv("Reply -> appointment rate", _pct(bookings, replied), is_pct=True)
    kv("Bounce rate", _pct(bounced, n_contacted), is_pct=True)
    kv("Opt-out rate", _pct(opted, n_contacted), is_pct=True)
    r += 1

    section("EMAILS SENT BY STEP")
    for s in range(1, 6):
        kv(f"Email {s} ({['day1','day3','day6','day10','day14'][s-1]})", sends_by_step[s])
    r += 1

    section("PERSONALIZATION LIFT")
    for flag, label in ((1, "Personalized"), (0, "Skipped / generic")):
        ct, rp = pers[flag]
        d.cell(r, 1, label)
        d.cell(r, 2, f"{rp}/{ct} replied")
        cell = d.cell(r, 3, _pct(rp, ct))
        cell.number_format = "0.0%"
        r += 1
    r += 1

    section("REPLY RATE BY NICHE")
    d.cell(r, 1, "Niche").font = HEADER_FONT
    d.cell(r, 1).fill = HEADER_FILL
    d.cell(r, 2, "Contacted").font = HEADER_FONT
    d.cell(r, 2).fill = HEADER_FILL
    d.cell(r, 3, "Replied").font = HEADER_FONT
    d.cell(r, 3).fill = HEADER_FILL
    d.cell(r, 4, "Reply rate").font = HEADER_FONT
    d.cell(r, 4).fill = HEADER_FILL
    r += 1
    for cat, (ct, rp) in sorted(niche.items(), key=lambda x: -x[1][0]):
        d.cell(r, 1, cat)
        d.cell(r, 2, ct)
        d.cell(r, 3, rp)
        cell = d.cell(r, 4, _pct(rp, ct))
        cell.number_format = "0.0%"
        r += 1

    d.column_dimensions["A"].width = 30
    for col in ("B", "C", "D"):
        d.column_dimensions[col].width = 16

    # ---------------- Log ----------------
    log = wb.create_sheet("Log")
    headers = ["Business", "Email", "Niche", "City", "Personalized",
               "Date sent", "Step", "Status", "Replied date"]
    log.append(headers)
    for cell in log[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in contacted:
        log.append([
            row["business"], row["email"], row["category"], row["city"],
            "Y" if row["personalized"] else ("N" if row["personalized"] == 0 else ""),
            row["first_contact_date"], row["current_step"], row["status"],
            row["replied_at"] or "",
        ])
    widths = [30, 34, 18, 16, 12, 12, 6, 12, 12]
    for i, w in enumerate(widths, 1):
        log.column_dimensions[get_column_letter(i)].width = w
    log.freeze_panes = "A2"

    wb.save(path)
    return path, n_contacted, total_sends


if __name__ == "__main__":
    p, nc, ts = build()
    print(f"wrote {p} | contacted={nc} sends={ts}")
